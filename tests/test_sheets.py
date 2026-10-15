"""Underwriting review of Excel lists: detection, cleaning, classes, formatted output, API."""

from __future__ import annotations

import io
from datetime import date

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.api.routes.sheets import superseded_lists
from app.insurance import sheets
from app.insurance.models import RequestDocument
from app.main import app

TODAY = date(2026, 10, 15)
HEADERS = ["م", "الاسم", "رقم الهوية", "تاريخ الميلاد", "الجنس", "صلة القرابة", "الفئة", "الجنسية"]
ROWS = [
    [1, "أحمد علي", 1012345678, date(1985, 3, 10), "ذكر", "موظف", "VIP", "سعودي"],
    [2, "سارة محمد", "1012345679", "1987-05-01", "F", "زوجة", "VIP", None],
    [3, "ليان أحمد", "١٠١٢٣٤٥٦٨٠", "01/01/2015", "أنثى", "ابنة", "VIP", None],
    [4, "John Smith", 2123456789, date(1970, 2, 2), "male", "employee", "Gold", "British"],
    [5, "خالد", "12345", date(1990, 1, 1), "ذكر", "موظف", "A", None],
    [6, "مكرر", 1012345678, date(1991, 1, 1), "ذكر", "موظف", "A", None],
    [7, "تاريخ هجري", 1011111111, "1405/01/01", "ذكر", "موظف", "B", None],
    [8, "كبير السن", 1022222222, date(1955, 1, 1), "ذكر", "موظف", "B", None],
]


def census_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(["كشف موظفي الشركة التجريبية"])
    ws.append([])
    ws.append(HEADERS)
    for row in ROWS:
        ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def fleet_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(["Plate No", "Sequence No", "Make", "Model", "Model Year", "Vehicle Value", "Usage"])
    ws.append(["ABC 1234", "123456789", "Toyota", "Hilux", 2022, "85,000", "تجاري"])
    ws.append(["XYZ 9876", "123456789", "Nissan", "Patrol", 2008, 60000, "خاص"])
    ws.append(["LMN 5555", "555555555", "Ford", "F150", 2035, -1, "نقل"])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _census() -> tuple[sheets.Sheet, dict[str, int | None]]:
    sheet = sheets.read_workbook(census_xlsx())
    return sheet, sheets.auto_mapping(sheets.PEOPLE, sheet.headers)


def test_header_row_and_columns_are_detected():
    sheet, mapping = _census()
    assert sheet.header_row == 3 and len(sheet.rows) == len(ROWS)
    assert mapping["name"] == 1 and mapping["id_number"] == 2 and mapping["birth_date"] == 3
    assert mapping["gender"] == 4 and mapping["relation"] == 5 and mapping["class"] == 6
    assert mapping["nationality"] == 7 and mapping["age"] is None


def test_census_is_cleaned_and_checked():
    sheet, mapping = _census()
    result = sheets.analyze(sheets.PEOPLE, sheet, sheets.Options(column_mapping=mapping), TODAY)
    rows = {r["_row"]: r for r in result.rows}
    assert rows[4]["id_number"] == "1012345678" and rows[4]["age"] == 41
    assert rows[5]["gender"] == "أنثى" and rows[5]["relation"] == "زوج/زوجة" and rows[5]["nationality"] == "سعودي"
    assert rows[6]["id_number"] == "1012345680" and rows[6]["relation"] == "ابن/ابنة"  # Arabic digits
    assert rows[7]["nationality"] == "British" and rows[7]["relation"] == "موظف"
    messages = [(i["row"], i["level"], i["message"]) for i in result.issues]
    assert any(r == 8 and lvl == "error" and "12345" in m for r, lvl, m in messages)
    assert any(r == 9 and "مكرر مع الصف 4" in m for r, _, m in messages)
    assert any(r == 11 and lvl == "warning" and "65" in m for r, lvl, m in messages)
    assert any(r is None and "هجري" in m for r, _, m in messages)
    assert 39 <= rows[10]["age"] <= 43  # 1405 AH ≈ 1985
    s = result.summary
    assert s["employees"] == 6 and s["dependents"] == 2 and s["saudis"] == 7 and s["over_65"] == 1
    assert {c["raw"]: c["count"] for c in result.categories} == {"VIP": 3, "A": 2, "B": 2, "Gold": 1}


def test_classes_are_mapped_and_filtered():
    sheet, mapping = _census()
    options = sheets.Options(column_mapping=mapping, class_mapping={"Gold": "A"}, include_classes=["VIP", "A"])
    result = sheets.analyze(sheets.PEOPLE, sheet, options, TODAY)
    assert {c["category"]: c["count"] for c in result.summary["by_category"]} == {"A": 3, "VIP": 3}
    assert result.summary["total_rows"] == 8 and result.summary["included_rows"] == 6
    clean = sheets.analyze(
        sheets.PEOPLE, sheet, sheets.Options(column_mapping=mapping, exclude_invalid=True), TODAY
    )
    assert clean.summary["included_rows"] == 6  # the invalid ID and the duplicate are left out


def test_formatted_workbook():
    sheet, mapping = _census()
    result = sheets.analyze(sheets.PEOPLE, sheet, sheets.Options(column_mapping=mapping), TODAY)
    content = sheets.build_workbook(sheets.PEOPLE, result, {"customer": "شركة تجريبية"}, split_by_class=True)
    wb = load_workbook(io.BytesIO(content))
    assert wb.sheetnames[:2] == ["الملخص", "الكشف"] and "الفئة VIP" in wb.sheetnames
    listing = wb["الكشف"]
    assert listing.sheet_view.rightToLeft and listing.cell(1, 2).value == "الاسم\nName"
    assert listing.max_row == len(ROWS) + 1
    assert wb["الفئة VIP"].max_row == 4


def test_fleet_profile():
    sheet = sheets.read_workbook(fleet_xlsx())
    mapping = sheets.auto_mapping(sheets.VEHICLES, sheet.headers)
    assert mapping["serial"] == 1 and mapping["year"] == 4 and mapping["value"] == 5 and mapping["usage"] == 6
    result = sheets.analyze(sheets.VEHICLES, sheet, sheets.Options(column_mapping=mapping), TODAY)
    messages = " | ".join(i["message"] for i in result.issues)
    assert "مكرر مع الصف 2" in messages and "2035" in messages and "القيمة غير صحيحة" in messages
    assert "أكثر من 15 سنة" in messages
    assert result.summary["total_value"] == 145000 - 1


def test_unreadable_file():
    try:
        sheets.read_workbook(b"PK\x03\x04 not a workbook")
    except sheets.SheetError as exc:
        assert "xlsx" in str(exc)
    else:
        raise AssertionError("expected SheetError")


async def test_underwriter_analyzes_targets_and_formats(db_session):
    with TestClient(app) as client:
        insurer = client.post("/insurance/insurers", json={"name": "شركة أ", "email": "q@a.example"}).json()
        created = client.post(
            "/insurance/portal/requests",
            json={
                "customer_type": "business",
                "full_name": "ريم",
                "company_name": "شركة تجريبية",
                "mobile": "0551234567",
                "insurance_type": "medical",
                "product_type": "medical_group",
                "details": {
                    "members_count": 8,
                    "network_class": "A",
                    "dependents_included": "نعم",
                    "start_date": "2027-01-01",
                },
                "consent": True,
            },
        ).json()
        upload = client.post(
            f"/insurance/portal/requests/{created['request_id']}/documents",
            files={"file": ("census.xlsx", census_xlsx())},
            data={"document_type": "employee_census"},
            headers={"X-Upload-Token": created["upload_token"]},
        )
        assert upload.status_code == 201 and upload.json()["review_status"] == "passed"
        doc_id = upload.json()["document_id"]

        lists = client.get(f"/insurance/requests/{created['request_id']}/lists").json()
        assert [x["document_id"] for x in lists] == [doc_id] and lists[0]["analysis"] is None

        first = client.post(f"/insurance/documents/{doc_id}/analysis").json()
        assert first["profile"] == "people" and first["header_row"] == 3
        assert first["issue_counts"]["error"] >= 2 and first["summary"]["total_rows"] == 8

        options = first["options"] | {
            "class_mapping": {"Gold": "A"},
            "include_classes": ["VIP", "A", "B"],
            "target_insurer_ids": [insurer["insurer_id"]],
            "exclude_invalid": True,
        }
        decided = client.post(f"/insurance/documents/{doc_id}/analysis", json=options).json()
        assert decided["summary"]["included_rows"] == 6 and decided["status"] == "draft"
        bad = client.post(f"/insurance/documents/{doc_id}/analysis", json={"target_insurer_ids": ["nope"]})
        assert bad.status_code == 422

        reopened = client.post(f"/insurance/documents/{doc_id}/analysis").json()  # saved decisions come back
        assert reopened["options"]["class_mapping"] == {"Gold": "A"}

        done = client.post(f"/insurance/analyses/{decided['analysis_id']}/finalize").json()
        assert done["status"] == "finalized" and done["output_document_id"]
        file = client.get(f"/insurance/documents/{done['output_document_id']}/download")
        assert file.status_code == 200
        wb = load_workbook(io.BytesIO(file.content))
        assert wb["الكشف"].max_row == 7
        summary = [[c.value for c in row] for row in wb["الملخص"].iter_rows()]
        assert ["الشركات المستهدفة", "شركة أ"] in [r[:2] for r in summary]

        again = client.post(f"/insurance/analyses/{decided['analysis_id']}/finalize").json()
        assert again["output_document_id"] != done["output_document_id"]
        review = client.get(f"/insurance/requests/{created['request_id']}/review").json()
        assert review["extra_documents"] == []  # formatted copies are BUD's, not the customer's

    assert await superseded_lists(db_session, created["request_id"]) == {doc_id}
    retired = await db_session.get(RequestDocument, done["output_document_id"])
    assert retired is not None and retired.review_status == "rejected"
