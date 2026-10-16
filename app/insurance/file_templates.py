"""Blank Excel templates for every list the system reads, built from the readers' own field
definitions so a template can never drift from what `sheets.read_workbook` understands.

    python -m app.cli export-templates     # writes app/static/dashboard/templates/*.xlsx
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.insurance import networks, sheets

OUTPUT_DIR = Path(__file__).resolve().parents[1] / "static" / "dashboard" / "templates"


@dataclass(frozen=True)
class Template:
    key: str
    filename: str
    title: str
    used_for: str
    profile: sheets.Profile
    fields: tuple[str, ...]
    rows: tuple[tuple[Any, ...], ...]
    notes: tuple[str, ...] = ()


ACCEPTED = {
    "gender": "ذكر / أنثى (أو M / F)",
    "relation": "موظف / زوج / زوجة / ابن / ابنة / والد / والدة (أو Employee / Spouse / Child / Parent)",
    "class": "VIP / A / B / C",
    "birth_date": "يوم/شهر/سنة ميلادي مثل 15/03/1990 (يُقبل الهجري ويُحوَّل)",
    "id_number": "10 أرقام تبدأ بـ 1 (مواطن) أو 2 (مقيم)",
    "passport": "حروف وأرقام بدون مسافات",
    "usage": "خاص / تجاري / نقل / تأجير",
    "year": "سنة ميلادية بأربعة أرقام",
    "value": "رقم بالريال بدون رموز",
    "sum_insured": "رقم بالريال بدون رموز",
    "network_class": "فئة واحدة أو عدة فئات مفصولة بفاصلة، مثل: VIP, A, B",
    "provider_type": "مستشفى / مستوصف / عيادة / صيدلية / مختبر / أشعة",
}

TEMPLATES = (
    Template(
        "medical_census",
        "كشف-الموظفين-والتابعين.xlsx",
        "كشف الموظفين والتابعين",
        "التأمين الطبي الجماعي — مستند «كشف الموظفين والتابعين»",
        sheets.PEOPLE,
        ("name", "id_number", "birth_date", "gender", "nationality", "relation", "marital_status", "class"),
        (
            ("محمد عبدالله العتيبي", "1012345678", "15/03/1985", "ذكر", "سعودي", "موظف", "متزوج", "A"),
            ("نورة سعد العتيبي", "1023456789", "02/07/1990", "أنثى", "سعودي", "زوجة", "متزوجة", "A"),
            ("سارة محمد العتيبي", "1034567890", "20/11/2015", "أنثى", "سعودي", "ابنة", "", "A"),
            ("Ahmed Khan", "2123456789", "10/01/1992", "M", "باكستاني", "Employee", "Single", "B"),
        ),
        ("كل تابع في صف مستقل تحت الموظف.", "اترك الفئة فارغة إذا لم تُحدد، وستُختار عند المراجعة."),
    ),
    Template(
        "members_list",
        "كشف-المؤمن-لهم.xlsx",
        "كشف المؤمَّن لهم",
        "الحماية والادخار والحوادث الشخصية وخيانة الأمانة — مستند «كشف المؤمَّن لهم»",
        sheets.PEOPLE,
        ("name", "id_number", "birth_date", "gender", "nationality", "sum_insured"),
        (
            ("خالد فهد القحطاني", "1045678901", "05/09/1980", "ذكر", "سعودي", 250000),
            ("Maria Santos", "2234567890", "12/04/1995", "F", "فلبينية", 150000),
        ),
        ("مبلغ التغطية لكل فرد (مثل الراتب السنوي × عدد مرات التغطية).",),
    ),
    Template(
        "travelers_list",
        "كشف-المسافرين.xlsx",
        "كشف المسافرين",
        "تأمين السفر — مستند «كشف المسافرين»",
        sheets.PEOPLE,
        ("name", "id_number", "passport", "birth_date", "gender", "nationality"),
        (
            ("عبدالرحمن سالم الدوسري", "1056789012", "U1234567", "18/06/1988", "ذكر", "سعودي"),
            ("ليلى سالم الدوسري", "1067890123", "U7654321", "25/12/2012", "أنثى", "سعودي"),
        ),
        ("أرفق صفحة بيانات الجواز لكل مسافر كمستند منفصل.",),
    ),
    Template(
        "vehicle_list",
        "كشف-المركبات.xlsx",
        "كشف المركبات",
        "تأمين أسطول المركبات — مستند «كشف المركبات»",
        sheets.VEHICLES,
        ("serial", "plate", "chassis", "make", "model", "year", "value", "usage"),
        (
            ("123456789", "أ ب ج 1234", "JTDBE32K123456789", "تويوتا", "كامري", 2023, 95000, "خاص"),
            ("987654321", "د هـ و 5678", "KMHCT41D987654321", "هيونداي", "أكسنت", 2022, 55000, "تأجير"),
            ("456789123", "س ص ط 9012", "WDB9066331S456789", "مرسيدس", "سبرنتر", 2021, 180000, "نقل"),
        ),
        ("الرقم التسلسلي من استمارة المركبة، أو رقم البطاقة الجمركية للمركبات الجديدة.",),
    ),
    Template(
        "assets",
        "جدول-الأصول-والمعدات.xlsx",
        "جدول الأصول والمعدات",
        "الممتلكات (جدول الأصول) والهندسي (كشف المعدات، جدول الكميات) والبحري (تقدير الشحنات)",
        sheets.GENERIC,
        ("description", "category", "serial", "location", "quantity", "value"),
        (
            ("المبنى الرئيسي", "مباني", "", "الرياض - حي العليا", 1, 4500000),
            ("أثاث ومعدات مكتبية", "محتويات", "", "الرياض - حي العليا", 1, 350000),
            ("رافعة شوكية", "معدات", "FL-2021-0088", "مستودع الدمام", 2, 120000),
        ),
        ("للشحنات: البند = نوع البضاعة، الكمية = عدد الشحنات المتوقع سنويًا، القيمة = إجمالي القيمة السنوية.",),
    ),
    Template(
        "providers",
        "شبكة-مقدمي-الخدمة.xlsx",
        "شبكة مقدمي الخدمة الطبية",
        "شركات التأمين — الشبكة الطبية لكل شركة (صفحة شركات التأمين)",
        networks.PROVIDERS,
        ("name", "provider_type", "city", "region", "network_class", "phone"),
        (
            ("مستشفى الحمادي", "مستشفى", "الرياض", "الوسطى", "VIP, A", "0114444444"),
            ("مجمع عيادات النخبة", "مستوصف", "جدة", "الغربية", "A, B, C", "0126666666"),
            ("صيدليات الدواء", "صيدلية", "الدمام", "الشرقية", "VIP, A, B, C", ""),
        ),
        ("إن كانت الشركة ترسل عمودًا لكل فئة (VIP، A، B…) بعلامة ✓، يقرأه النظام أيضًا.",),
    ),
)

HEADER_FILL = PatternFill("solid", fgColor="0A0A0B")
REQUIRED_FILL = PatternFill("solid", fgColor="2D5BFF")


def _field(profile: sheets.Profile, key: str) -> sheets.FieldDef:
    return next(f for f in profile.fields if f.key == key)


def build(template: Template) -> Workbook:
    wb = Workbook()
    data = wb.worksheets[0]
    data.title = "البيانات"
    data.sheet_view.rightToLeft = True
    fields = [_field(template.profile, k) for k in template.fields]
    data.append([f.label for f in fields])
    for cell, f in zip(data[1], fields, strict=True):
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = REQUIRED_FILL if f.required else HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in template.rows:
        data.append(list(row))
    for i, f in enumerate(fields, start=1):
        data.column_dimensions[get_column_letter(i)].width = max(14, len(f.label) + 6)
    data.freeze_panes = "A2"

    guide = wb.create_sheet("تعليمات")
    guide.sheet_view.rightToLeft = True
    guide.append([template.title])
    guide["A1"].font = Font(bold=True, size=14)
    guide.append([f"يُستخدم في: {template.used_for}"])
    guide.append(["امسح صفوف الأمثلة واكتب بياناتك تحت العناوين نفسها. لا تغيّر أسماء الأعمدة."])
    guide.append(["العناوين الزرقاء إلزامية، والسوداء اختيارية."])
    for note in template.notes:
        guide.append([note])
    guide.append([])
    guide.append(["العمود", "بالإنجليزية", "إلزامي", "القيم المقبولة"])
    for cell in guide[guide.max_row]:
        cell.font = Font(bold=True)
    for f in fields:
        guide.append([f.label, f.english, "نعم" if f.required else "لا", ACCEPTED.get(f.key, "")])
    for column, width in zip("ABCD", (24, 18, 10, 70), strict=True):
        guide.column_dimensions[column].width = width
    return wb


def export(directory: Path = OUTPUT_DIR) -> list[Path]:
    directory.mkdir(parents=True, exist_ok=True)
    paths = []
    for template in TEMPLATES:
        path = directory / template.filename
        build(template).save(path)
        paths.append(path)
    return paths
