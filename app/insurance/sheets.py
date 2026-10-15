"""Underwriting lists in Excel: census of members, vehicle fleets, asset/equipment schedules.

`read_workbook` finds the header row and maps columns to standard fields (Arabic or
English headings); `analyze` cleans every row (IDs, dates, gender, relation, numbers),
reports problems by row, and summarizes the risk (classes, age bands, values);
`build_workbook` writes the clean, insurer-ready file. All rules, no model.
"""

from __future__ import annotations

import io
import re
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

MAX_ROWS = 20_000
HEADER_SCAN = 15


@dataclass(frozen=True)
class FieldDef:
    key: str
    label: str  # Arabic heading in the output
    english: str  # English heading in the output (insurers' systems often expect it)
    synonyms: tuple[str, ...]
    required: bool = False


@dataclass(frozen=True)
class Profile:
    key: str
    label: str
    fields: tuple[FieldDef, ...]
    category_field: str  # the column grouped into classes/categories
    category_label: str
    value_field: str | None = None
    default_categories: tuple[str, ...] = ()
    documents: tuple[str, ...] = field(default=())


PEOPLE = Profile(
    "people",
    "كشف أفراد",
    (
        FieldDef(
            "name",
            "الاسم",
            "Name",
            (
                "name",
                "full name",
                "employee name",
                "member name",
                "الاسم",
                "اسم الموظف",
                "الاسم الكامل",
                "اسم المؤمن له",
                "اسم العضو",
            ),
            True,
        ),
        FieldDef(
            "id_number",
            "رقم الهوية/الإقامة",
            "ID / Iqama",
            (
                "id",
                "id no",
                "id number",
                "iqama",
                "iqama no",
                "national id",
                "رقم الهوية",
                "الهوية",
                "رقم الاقامة",
                "الاقامة",
                "هوية",
                "اقامة",
                "السجل المدني",
            ),
            True,
        ),
        FieldDef(
            "birth_date",
            "تاريخ الميلاد",
            "Date of Birth",
            ("dob", "date of birth", "birth date", "birthdate", "تاريخ الميلاد", "الميلاد"),
        ),
        FieldDef("age", "العمر", "Age", ("age", "العمر", "السن")),
        FieldDef("gender", "الجنس", "Gender", ("gender", "sex", "الجنس", "النوع"), True),
        FieldDef("nationality", "الجنسية", "Nationality", ("nationality", "الجنسيه", "الجنسية")),
        FieldDef(
            "relation",
            "صلة القرابة",
            "Relation",
            (
                "relation",
                "relationship",
                "member type",
                "dependency",
                "صله القرابه",
                "العلاقه",
                "القرابه",
                "نوع العضو",
                "صفه",
            ),
        ),
        FieldDef(
            "marital_status",
            "الحالة الاجتماعية",
            "Marital Status",
            ("marital status", "marital", "الحاله الاجتماعيه"),
        ),
        FieldDef(
            "class",
            "الفئة",
            "Class",
            ("class", "category", "network", "grade", "plan", "الفئه", "الدرجه", "الشبكه", "فئه", "فئة التامين"),
        ),
        FieldDef(
            "sum_insured",
            "مبلغ التغطية",
            "Sum Insured",
            ("sum insured", "coverage", "salary", "basic salary", "مبلغ التغطيه", "مبلغ التامين", "الراتب"),
        ),
    ),
    category_field="class",
    category_label="الفئة",
    value_field="sum_insured",
    default_categories=("VIP", "A", "B", "C"),
    documents=("employee_census", "members_list", "insured_list"),
)
VEHICLES = Profile(
    "vehicles",
    "كشف مركبات",
    (
        FieldDef("plate", "رقم اللوحة", "Plate No", ("plate", "plate no", "plate number", "رقم اللوحه", "اللوحه")),
        FieldDef(
            "serial",
            "الرقم التسلسلي",
            "Sequence No",
            (
                "sequence",
                "sequence no",
                "serial",
                "serial no",
                "custom card",
                "الرقم التسلسلي",
                "التسلسلي",
                "رقم التسلسل",
            ),
            True,
        ),
        FieldDef(
            "chassis", "رقم الهيكل", "Chassis / VIN", ("chassis", "chassis no", "vin", "رقم الهيكل", "الهيكل")
        ),
        FieldDef(
            "make",
            "الشركة المصنعة",
            "Make",
            ("make", "manufacturer", "brand", "الشركه المصنعه", "الماركه", "الصانع", "النوع"),
            True,
        ),
        FieldDef("model", "الطراز", "Model", ("model", "الطراز", "الموديل"), True),
        FieldDef(
            "year",
            "سنة الصنع",
            "Year",
            ("year", "model year", "manufacturing year", "سنه الصنع", "سنة الصنع"),
            True,
        ),
        FieldDef(
            "value",
            "القيمة (ر.س)",
            "Value (SAR)",
            (
                "value",
                "sum insured",
                "market value",
                "vehicle value",
                "القيمه",
                "قيمه المركبه",
                "القيمه التقديريه",
                "مبلغ التامين",
            ),
            True,
        ),
        FieldDef(
            "usage",
            "الاستخدام",
            "Usage",
            ("usage", "use", "category", "type", "الاستخدام", "نوع الاستخدام", "الفئه", "نوع المركبه"),
        ),
    ),
    category_field="usage",
    category_label="الاستخدام",
    value_field="value",
    default_categories=("خاص", "تجاري", "نقل", "تأجير"),
    documents=("vehicle_list",),
)
GENERIC = Profile(
    "generic",
    "جدول بنود",
    (
        FieldDef(
            "description",
            "البند",
            "Item",
            ("description", "item", "asset", "equipment", "البند", "الوصف", "الصنف", "الاصل", "المعده"),
            True,
        ),
        FieldDef("serial", "الرقم التسلسلي", "Serial No", ("serial", "serial no", "الرقم التسلسلي")),
        FieldDef("location", "الموقع", "Location", ("location", "site", "الموقع", "المدينه")),
        FieldDef("quantity", "الكمية", "Qty", ("qty", "quantity", "الكميه", "العدد")),
        FieldDef(
            "value",
            "القيمة (ر.س)",
            "Value (SAR)",
            ("value", "amount", "sum insured", "القيمه", "المبلغ", "مبلغ التامين"),
            True,
        ),
        FieldDef("category", "الفئة", "Category", ("category", "type", "class", "الفئه", "النوع")),
    ),
    category_field="category",
    category_label="الفئة",
    value_field="value",
)
PROFILES = {p.key: p for p in (PEOPLE, VEHICLES, GENERIC)}


def profile_for(document_type: str) -> Profile:
    return next((p for p in PROFILES.values() if document_type in p.documents), GENERIC)


# --- Normalizing text ----------------------------------------------------------------------

_ARABIC_DIGITS = str.maketrans("٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹", "01234567890123456789")
_DIACRITICS = re.compile(r"[ً-ْـ]")


def norm(text: Any) -> str:
    """Lowercase, unify Arabic letter forms, drop diacritics and punctuation, for matching."""
    s = _DIACRITICS.sub("", str(text or "")).translate(_ARABIC_DIGITS).strip().lower()
    s = re.sub("[أإآ]", "ا", s).replace("ة", "ه").replace("ى", "ي")
    s = re.sub(r"[_\-./:()\[\]#*]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def _match(header: str, f: FieldDef) -> int:
    """2 = exact synonym, 1 = header contains a synonym of 3+ letters, 0 = no match."""
    h = norm(header)
    if not h:
        return 0
    synonyms = [norm(s) for s in f.synonyms]
    if h in synonyms:
        return 2
    return 1 if any(len(s) >= 3 and re.search(rf"(^|\s){re.escape(s)}($|\s)", h) for s in synonyms) else 0


# --- Reading -------------------------------------------------------------------------------


@dataclass
class Sheet:
    sheet_names: list[str]
    sheet_name: str
    header_row: int  # 1-based, as Excel shows it
    headers: list[str]
    rows: list[tuple[int, list[Any]]]  # (Excel row number, cells)
    truncated: bool


class SheetError(Exception):
    pass


def read_workbook(content: bytes, sheet_name: str | None = None) -> Sheet:
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises many types for a bad file
        raise SheetError("تعذّرت قراءة ملف Excel. تأكد أنه بصيغة xlsx وغير محمي بكلمة مرور.") from exc
    names = list(wb.sheetnames)
    if sheet_name and sheet_name not in names:
        raise SheetError(f"لا توجد ورقة باسم «{sheet_name}» في الملف.")
    grids: dict[str, list[list[Any]]] = {}
    for name in [sheet_name] if sheet_name else names:
        grid = []
        for values in wb[name].iter_rows(values_only=True):
            grid.append(list(values))
            if len(grid) > MAX_ROWS + HEADER_SCAN:
                break
        grids[name] = grid
    wb.close()
    # Without a choice, take the sheet with the most filled rows.
    chosen = sheet_name or max(
        grids, key=lambda n: sum(1 for r in grids[n] if any(c not in (None, "") for c in r))
    )
    grid = grids[chosen]
    if not any(any(c not in (None, "") for c in r) for r in grid):
        raise SheetError("الورقة فارغة.")

    all_fields = [f for p in PROFILES.values() for f in p.fields]

    def score(row: list[Any]) -> int:
        return sum(1 for c in row if isinstance(c, str) and any(_match(c, f) for f in all_fields))

    candidates = [
        (score(r), -i, i) for i, r in enumerate(grid[:HEADER_SCAN]) if any(c not in (None, "") for c in r)
    ]
    _, _, header_index = max(candidates)
    width = max(len(r) for r in grid)
    headers = [
        str(c).strip() if c not in (None, "") else f"عمود {i + 1}"
        for i, c in enumerate(grid[header_index] + [None] * (width - len(grid[header_index])))
    ]
    rows = [
        (header_index + 2 + i, r + [None] * (width - len(r)))
        for i, r in enumerate(grid[header_index + 1 :])
        if any(c not in (None, "") for c in r)
    ]
    return Sheet(names, chosen, header_index + 1, headers, rows[:MAX_ROWS], len(rows) > MAX_ROWS)


def auto_mapping(profile: Profile, headers: list[str]) -> dict[str, int | None]:
    """Best column for each field; each column used once, exact names first."""
    taken: set[int] = set()
    mapping: dict[str, int | None] = {f.key: None for f in profile.fields}
    for strength in (2, 1):
        for f in profile.fields:
            if mapping[f.key] is not None:
                continue
            for i, h in enumerate(headers):
                if i not in taken and _match(h, f) == strength:
                    mapping[f.key] = i
                    taken.add(i)
                    break
    return mapping


# --- Cleaning values -----------------------------------------------------------------------


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).translate(_ARABIC_DIGITS).strip()


def _number(value: Any) -> float | None:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    s = _text(value).replace(",", "").replace("٬", "").replace("ر.س", "").replace("SAR", "").strip()
    try:
        return float(s) if s else None
    except ValueError:
        return None


def _hijri_to_gregorian_year(year: int) -> int:
    return round(year * 0.970229 + 621.5643)


def _date(value: Any) -> tuple[date | None, bool]:
    """(date, was_hijri). Hijri dates are converted approximately (enough for age bands)."""
    if isinstance(value, datetime):
        return value.date(), False
    if isinstance(value, date):
        return value, False
    s = _text(value)
    m = re.match(r"^(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})", s) or None
    parts: tuple[int, int, int] | None = None
    if m:
        parts = (int(m[1]), int(m[2]), int(m[3]))
    else:
        m = re.match(r"^(\d{1,2})[-/.](\d{1,2})[-/.](\d{4})", s)
        if m:
            parts = (int(m[3]), int(m[2]), int(m[1]))
    if parts is None:
        return None, False
    year, month, day = parts
    hijri = 1300 <= year <= 1500
    if hijri:
        year = _hijri_to_gregorian_year(year)
    try:
        return date(year, min(max(month, 1), 12), min(max(day, 1), 28) if hijri else day), hijri
    except ValueError:
        return None, False


GENDERS: dict[str, tuple[str, ...]] = {
    "ذكر": ("m", "male", "ذكر", "رجل", "1"),
    "أنثى": ("f", "female", "انثي", "امراه", "2"),
}
RELATIONS: dict[str, tuple[str, ...]] = {
    "موظف": ("employee", "self", "principal", "main", "member", "موظف", "الموظف", "رئيسي", "نفسه", "صاحب العمل"),
    "زوج/زوجة": ("spouse", "wife", "husband", "زوج", "زوجه"),
    "ابن/ابنة": ("child", "son", "daughter", "kid", "ابن", "ابنه", "بنت", "ولد", "طفل"),
    "والد/والدة": ("parent", "father", "mother", "اب", "ام", "والد", "والده"),
}


def _lookup(value: Any, table: dict[str, tuple[str, ...]]) -> str | None:
    v = norm(value)
    return next((label for label, words in table.items() if v in {norm(w) for w in words}), None)


AGE_BANDS = [(0, 17), (18, 29), (30, 39), (40, 49), (50, 59), (60, 64), (65, 200)]
YEAR_BANDS = [(0, 3, "حتى 3 سنوات"), (4, 7, "4 - 7 سنوات"), (8, 12, "8 - 12 سنة"), (13, 200, "أكثر من 12 سنة")]


def _band(age: int) -> str:
    low, high = next(b for b in AGE_BANDS if b[0] <= age <= b[1])
    return f"{low}+" if high >= 200 else f"{low}-{high}"


# --- Analysis ------------------------------------------------------------------------------


@dataclass
class Options:
    column_mapping: dict[str, int | None]
    class_mapping: dict[str, str] = field(default_factory=dict)  # raw category → target category
    default_class: str | None = None  # for rows with no category
    include_classes: list[str] | None = None  # None = every class
    exclude_invalid: bool = False  # leave rows with errors out of the output


@dataclass
class Result:
    rows: list[dict[str, Any]]  # clean rows (with "_row" and "_errors")
    issues: list[dict[str, Any]]
    summary: dict[str, Any]
    categories: list[dict[str, Any]]


def analyze(profile: Profile, sheet: Sheet, options: Options, today: date | None = None) -> Result:
    today = today or date.today()
    mapping = options.column_mapping
    issues: list[dict[str, Any]] = []
    clean: list[dict[str, Any]] = []
    raw_categories: Counter[str] = Counter()
    seen: dict[str, dict[str, int]] = {"id_number": {}, "serial": {}, "chassis": {}}
    hijri_rows = 0

    def issue(row: int, field_key: str, message: str, level: str = "error") -> None:
        issues.append({"row": row, "field": field_key, "message": message, "level": level})

    for excel_row, cells in sheet.rows:
        raw = {k: (cells[i] if i is not None and i < len(cells) else None) for k, i in mapping.items()}
        out: dict[str, Any] = {"_row": excel_row}
        before = len(issues)
        for f in profile.fields:
            if f.required and mapping.get(f.key) is not None and _text(raw.get(f.key)) == "":
                issue(excel_row, f.key, f"«{f.label}» فارغ")

        if profile is PEOPLE:
            out["name"] = re.sub(r"\s+", " ", _text(raw["name"]))
            digits = re.sub(r"\D", "", _text(raw["id_number"]))
            out["id_number"] = digits
            if digits and not re.fullmatch(r"[12]\d{9}", digits):
                issue(excel_row, "id_number", f"رقم الهوية {digits} غير صالح (10 أرقام تبدأ بـ 1 أو 2)")
            born, hijri = _date(raw["birth_date"])
            hijri_rows += hijri
            age_value = _number(raw["age"])
            if born:
                age = today.year - born.year - ((today.month, today.day) < (born.month, born.day))
                out["birth_date"] = born
            elif age_value is not None:
                age = int(age_value)
                out["birth_date"] = None
            else:
                age = None
                out["birth_date"] = None
                if mapping.get("birth_date") is not None or mapping.get("age") is not None:
                    issue(excel_row, "birth_date", "تاريخ الميلاد أو العمر غير مقروء")
            if age is not None and not 0 <= age <= 100:
                issue(excel_row, "age", f"العمر {age} غير منطقي")
                age = None
            out["age"] = age
            gender = _lookup(raw["gender"], GENDERS)
            out["gender"] = gender or ""
            if raw["gender"] not in (None, "") and gender is None:
                issue(excel_row, "gender", f"الجنس «{_text(raw['gender'])}» غير معروف")
            relation = _lookup(raw["relation"], RELATIONS)
            out["relation"] = relation or "موظف"
            if raw["relation"] not in (None, "") and relation is None:
                issue(excel_row, "relation", f"صلة القرابة «{_text(raw['relation'])}» غير معروفة", "warning")
            nationality = _text(raw["nationality"])
            if norm(nationality) in {"سعودي", "سعوديه", "saudi", "ksa", "saudi arabia"}:
                nationality = "سعودي"
            elif not nationality and digits[:1] in ("1", "2"):
                nationality = "سعودي" if digits[0] == "1" else "غير سعودي"
            out["nationality"] = nationality
            out["marital_status"] = _text(raw["marital_status"])
            out["sum_insured"] = _number(raw["sum_insured"])
            if age is not None and age >= 65:
                issue(
                    excel_row,
                    "age",
                    f"عمره {age} سنة: كثير من الشركات تستثني أو تسعّر من 65 فأكثر بشكل خاص",
                    "warning",
                )
            if out["relation"] == "ابن/ابنة" and age is not None and age > 25:
                issue(excel_row, "relation", f"تابع (ابن/ابنة) عمره {age} سنة، تحقق من الأهلية", "warning")
            key_fields: tuple[str, ...] = ("id_number",)
        elif profile is VEHICLES:
            for k in ("plate", "serial", "chassis", "make", "model"):
                out[k] = re.sub(r"\s+", " ", _text(raw[k]))
            out["serial"] = re.sub(r"\D", "", out["serial"]) or out["serial"]
            year = _number(raw["year"])
            out["year"] = int(year) if year else None
            if year is not None and not 1980 <= year <= today.year + 1:
                issue(excel_row, "year", f"سنة الصنع {int(year)} غير منطقية")
            value = _number(raw["value"])
            out["value"] = value
            if raw["value"] not in (None, "") and (value is None or value <= 0):
                issue(excel_row, "value", "القيمة غير صحيحة")
            out["usage"] = _text(raw["usage"])
            if out["year"] and today.year - out["year"] > 15:
                issue(excel_row, "year", "عمر المركبة أكثر من 15 سنة", "warning")
            key_fields = ("serial", "chassis")
        else:
            out["description"] = _text(raw["description"])
            out["serial"] = _text(raw["serial"])
            out["location"] = _text(raw["location"])
            quantity = _number(raw["quantity"])
            out["quantity"] = quantity
            value = _number(raw["value"])
            out["value"] = value
            if raw["value"] not in (None, "") and value is None:
                issue(excel_row, "value", "القيمة غير مقروءة")
            out["category"] = _text(raw["category"])
            key_fields = ("serial",)

        for k in key_fields:
            v = out.get(k)
            if v and k in seen:
                if v in seen[k]:
                    issue(excel_row, k, f"مكرر مع الصف {seen[k][v]}")
                else:
                    seen[k][v] = excel_row

        raw_category = _text(out.get(profile.category_field)) if profile is not PEOPLE else _text(raw["class"])
        raw_categories[raw_category] += 1
        target = options.class_mapping.get(raw_category) or raw_category or (options.default_class or "")
        if not target:
            issue(excel_row, profile.category_field, f"«{profile.category_label}» غير محدد", "warning")
        out[profile.category_field] = target
        out["_errors"] = any(i["level"] == "error" for i in issues[before:])
        clean.append(out)

    if hijri_rows:
        issues.append(
            {
                "row": None,
                "field": "birth_date",
                "level": "warning",
                "message": f"{hijri_rows} تاريخ ميلاد هجري حُوِّل للميلادي تقريبيًا لحساب العمر",
            }
        )

    included = [
        r
        for r in clean
        if (options.include_classes is None or r[profile.category_field] in options.include_classes)
        and not (options.exclude_invalid and r["_errors"])
    ]
    categories = [
        {"raw": raw, "count": n, "mapped": options.class_mapping.get(raw) or raw or (options.default_class or "")}
        for raw, n in raw_categories.most_common()
    ]
    return Result(included, issues, _summary(profile, clean, included, today), categories)


def _summary(
    profile: Profile, clean: list[dict[str, Any]], rows: list[dict[str, Any]], today: date
) -> dict[str, Any]:
    cat = profile.category_field
    summary: dict[str, Any] = {
        "total_rows": len(clean),
        "included_rows": len(rows),
        "rows_with_errors": sum(1 for r in clean if r["_errors"]),
    }
    by_category: dict[str, dict[str, Any]] = {}
    for r in rows:
        entry = by_category.setdefault(r[cat] or "غير محدد", {"count": 0, "value": 0.0})
        entry["count"] += 1
        if profile.value_field and r.get(profile.value_field):
            entry["value"] += r[profile.value_field]
    if profile is PEOPLE:
        for r in rows:
            entry = by_category[r[cat] or "غير محدد"]
            key = "employees" if r["relation"] == "موظف" else "dependents"
            entry[key] = entry.get(key, 0) + 1
        ages = [r["age"] for r in rows if r["age"] is not None]
        bands = [_band(a[0]) for a in AGE_BANDS]
        table = {b: {"ذكر": 0, "أنثى": 0} for b in bands}
        for r in rows:
            if r["age"] is not None and r["gender"] in ("ذكر", "أنثى"):
                table[_band(r["age"])][r["gender"]] += 1
        summary |= {
            "employees": sum(1 for r in rows if r["relation"] == "موظف"),
            "dependents": sum(1 for r in rows if r["relation"] != "موظف"),
            "by_relation": dict(Counter(r["relation"] for r in rows)),
            "by_gender": dict(Counter(r["gender"] or "غير محدد" for r in rows)),
            "saudis": sum(1 for r in rows if r["nationality"] == "سعودي"),
            "non_saudis": sum(1 for r in rows if r["nationality"] and r["nationality"] != "سعودي"),
            "average_age": round(sum(ages) / len(ages), 1) if ages else None,
            "age_bands": [{"band": b, **table[b]} for b in bands],
            "over_65": sum(1 for a in ages if a >= 65),
        }
    elif profile is VEHICLES:
        years = [today.year - r["year"] for r in rows if r.get("year")]
        summary |= {
            "total_value": sum(r["value"] or 0 for r in rows),
            "average_age": round(sum(years) / len(years), 1) if years else None,
            "by_make": dict(Counter(r["make"] or "غير محدد" for r in rows).most_common(10)),
            "by_age": {label: sum(1 for y in years if lo <= y <= hi) for lo, hi, label in YEAR_BANDS},
        }
    else:
        summary["total_value"] = sum(r["value"] or 0 for r in rows)
    summary["by_category"] = [{"category": k, **v} for k, v in sorted(by_category.items())]
    return summary


# --- Output --------------------------------------------------------------------------------

_HEAD = PatternFill("solid", start_color="1F3A8A")
_SUB = PatternFill("solid", start_color="E3EAFF")
_THIN = Side(style="thin", color="C9CED6")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _table(ws: Any, start_row: int, headers: list[str], rows: list[list[Any]]) -> int:
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=start_row, column=j, value=h)
        c.font, c.fill, c.border = Font(bold=True, color="FFFFFF"), _HEAD, _BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, row in enumerate(rows, start_row + 1):
        for j, v in enumerate(row, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.border = _BORDER
            if isinstance(v, float):
                c.number_format = "#,##0.00"
            elif isinstance(v, date):
                c.number_format = "yyyy-mm-dd"
    return start_row + len(rows) + 2


def _widths(ws: Any, columns: int, width: int = 18) -> None:
    for j in range(1, columns + 1):
        ws.column_dimensions[get_column_letter(j)].width = width


def build_workbook(
    profile: Profile,
    result: Result,
    title: dict[str, str],
    split_by_class: bool = False,
) -> bytes:
    wb = Workbook()
    summary_ws = wb.active
    assert summary_ws is not None
    summary_ws.title = "الملخص"
    summary_ws.sheet_view.rightToLeft = True
    summary_ws["A1"] = f"{profile.label}: {title.get('customer', '')}"
    summary_ws["A1"].font = Font(bold=True, size=14)
    line = 2
    for k, v in title.items():
        if k != "customer":
            summary_ws.cell(row=line, column=1, value=k).font = Font(bold=True)
            summary_ws.cell(row=line, column=2, value=v)
            line += 1
    s = result.summary
    facts = [("عدد السجلات", s["included_rows"])]
    if profile is PEOPLE:
        facts += [
            ("الموظفون", s["employees"]),
            ("التابعون", s["dependents"]),
            ("سعوديون", s["saudis"]),
            ("غير سعوديين", s["non_saudis"]),
            ("متوسط العمر", s["average_age"]),
        ]
    elif "total_value" in s:
        facts += [("إجمالي القيمة (ر.س)", float(s["total_value"]))]
    line = _table(summary_ws, line + 1, ["البيان", "القيمة"], [list(f) for f in facts])
    category_head = [profile.category_label, "العدد"] + (
        ["موظفون", "تابعون"] if profile is PEOPLE else ["إجمالي القيمة (ر.س)"]
    )
    category_rows = [
        [c["category"], c["count"]]
        + ([c.get("employees", 0), c.get("dependents", 0)] if profile is PEOPLE else [float(c["value"])])
        for c in s["by_category"]
    ]
    line = _table(summary_ws, line, category_head, category_rows)
    if profile is PEOPLE:
        _table(
            summary_ws,
            line,
            ["الفئة العمرية", "ذكور", "إناث"],
            [[b["band"], b["ذكر"], b["أنثى"]] for b in s["age_bands"]],
        )
    _widths(summary_ws, 4, 24)

    fields = list(profile.fields)
    if profile is PEOPLE:
        fields = [
            f for f in fields if f.key != "marital_status" or any(r.get("marital_status") for r in result.rows)
        ]
    fields = [f for f in fields if any(r.get(f.key) not in (None, "") for r in result.rows) or f.required]
    headers = ["#"] + [f"{f.label}\n{f.english}" for f in fields]

    def write_list(name: str, rows: list[dict[str, Any]]) -> None:
        ws = wb.create_sheet(name[:31])
        ws.sheet_view.rightToLeft = True
        _table(ws, 1, headers, [[i] + [r.get(f.key) for f in fields] for i, r in enumerate(rows, 1)])
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
        ws.row_dimensions[1].height = 32
        _widths(ws, len(headers))
        ws.column_dimensions["A"].width = 6

    write_list("الكشف", result.rows)
    if split_by_class:
        for c in s["by_category"]:
            write_list(
                f"{profile.category_label} {c['category']}",
                [r for r in result.rows if (r[profile.category_field] or "غير محدد") == c["category"]],
            )
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
